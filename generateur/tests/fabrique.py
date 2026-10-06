"""Fabrique de modèles Word de TEST, générés par du code (aucun modèle réel dans le dépôt).

    document("Bonjour {{ souscripteur.nom }}", [["Nom", "Date"], ["{{ a.nom }}", "…"]], entete="…")

Un texte donne un paragraphe ; une liste de textes, un paragraphe en plusieurs morceaux (runs),
comme Word en produit ; une liste de lignes (listes de cellules), un tableau.
"""

import io
from collections.abc import Sequence

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

Paragraphe = str | Sequence[str]
Tableau = Sequence[Sequence[str]]


def _remplir(paragraphe, contenu: Paragraphe) -> None:  # type: ignore[no-untyped-def]
    if isinstance(contenu, str):
        paragraphe.add_run(contenu)
    else:
        for i, morceau in enumerate(contenu):
            run = paragraphe.add_run(morceau)
            run.bold = i % 2 == 1  # mises en forme différentes : Word ne fusionnerait pas ces morceaux


def document(*blocs: Paragraphe | Tableau, entete: str | None = None, police: str | None = None) -> bytes:
    doc = Document()
    if police:
        doc.styles["Normal"].font.name = police
        doc.styles["Normal"].font.size = Pt(11)
    for bloc in blocs:
        if isinstance(bloc, (str, tuple)) or (isinstance(bloc, list) and bloc and isinstance(bloc[0], str)):
            _remplir(doc.add_paragraph(), bloc)  # type: ignore[arg-type]
        else:
            lignes = list(bloc)  # type: ignore[arg-type]
            tableau = doc.add_table(rows=len(lignes), cols=max(len(l) for l in lignes))
            tableau.style = "Table Grid"
            for i, ligne in enumerate(lignes):
                for j, texte in enumerate(ligne):
                    tableau.cell(i, j).text = texte
    if entete is not None:
        doc.sections[0].header.paragraphs[0].text = entete
    sortie = io.BytesIO()
    doc.save(sortie)
    return sortie.getvalue()


def avec_champ_word(texte_avant: str, instruction: str) -> bytes:
    """Paragraphe contenant un champ Word (Ctrl+F9) dont l'instruction est `instruction`."""
    doc = Document()
    paragraphe = doc.add_paragraph(texte_avant)
    for genre, texte in (("begin", None), (None, instruction), ("separate", None), (None, "résultat"), ("end", None)):
        run = paragraphe.add_run()
        if genre:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), genre)
        elif texte == instruction:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = texte
        else:
            element = OxmlElement("w:t")
            element.text = texte
        run._r.append(element)
    sortie = io.BytesIO()
    doc.save(sortie)
    return sortie.getvalue()


def avec_texte_supprime(avant: str, supprime: str, apres: str) -> bytes:
    """Paragraphe « avant + apres », avec `supprime` entre les deux en suivi des modifications
    (texte barré, w:delText) : Word l'affiche barré, docxtpl ne sait pas recoller autour."""
    doc = Document()
    paragraphe = doc.add_paragraph()
    paragraphe.add_run(avant)
    suppression = OxmlElement("w:del")
    suppression.set(qn("w:id"), "1")
    suppression.set(qn("w:author"), "Test")
    run = OxmlElement("w:r")
    texte = OxmlElement("w:delText")
    texte.text = supprime
    run.append(texte)
    suppression.append(run)
    paragraphe._p.append(suppression)
    paragraphe.add_run(apres)
    sortie = io.BytesIO()
    doc.save(sortie)
    return sortie.getvalue()


def texte_docx(contenu: bytes) -> str:
    """Texte d'un .docx, un paragraphe par ligne (tableaux compris, dans l'ordre du document)."""
    corps = Document(io.BytesIO(contenu)).element.body
    lignes = []
    for p in corps.iter(qn("w:p")):
        lignes.append("".join(t.text or "" for t in p.iter(qn("w:t"))))
    return "\n".join(lignes)


# Image PNG de 1 × 1 pixel (gris), pour les tests d'images.
PNG_1PX = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108000000003a7e9b55"
    "0000000a49444154789c636800000082008177cd72b60000000049454e44ae426082"
)

__all__ = ["PNG_1PX", "avec_champ_word", "avec_texte_supprime", "document", "texte_docx"]
