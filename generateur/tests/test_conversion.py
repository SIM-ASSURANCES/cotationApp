"""Tests de la conversion. Ils lancent le vrai LibreOffice : ignorés s'il est absent."""

import os
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader

from app.conversion import _PROFILS, ConversionDelaiDepasse, ConversionEchouee, convertir_en_pdf

CHEMINS_POSSIBLES = [
    Path(os.environ.get("LIBREOFFICE_REEL", "")),
    Path(r"C:\Program Files\LibreOffice\program\soffice.com"),
    Path("/usr/bin/soffice"),
]
SOFFICE = next((c for c in CHEMINS_POSSIBLES if c.is_file()), None)

pytestmark = pytest.mark.skipif(SOFFICE is None, reason="LibreOffice absent")


@pytest.fixture
def docx_de_test(tmp_path: Path) -> Path:
    document = Document()
    document.add_paragraph("Document de test inventé : conversion.")
    chemin = tmp_path / "essai.docx"
    document.save(chemin)
    return chemin


def test_convertit_un_docx_en_pdf(docx_de_test: Path, tmp_path: Path) -> None:
    assert SOFFICE is not None
    pdf = convertir_en_pdf(
        docx_de_test, tmp_path / "sortie", executable=SOFFICE, profil=tmp_path / "profil", delai_s=120
    )

    assert pdf.read_bytes().startswith(b"%PDF")
    assert "conversion" in PdfReader(pdf).pages[0].extract_text()


def test_interrompt_une_conversion_trop_longue(docx_de_test: Path, tmp_path: Path) -> None:
    assert SOFFICE is not None
    # 10 ms : le démarrage de LibreOffice prend bien plus, donc le délai est dépassé.
    with pytest.raises(ConversionDelaiDepasse):
        convertir_en_pdf(
            docx_de_test, tmp_path / "sortie", executable=SOFFICE, profil=tmp_path / "profil", delai_s=0.01
        )
    assert not (tmp_path / "sortie" / "essai.pdf").exists()


def test_fichier_inexistant_est_une_echec_sans_pdf(tmp_path: Path) -> None:
    assert SOFFICE is not None
    with pytest.raises(ConversionEchouee):
        convertir_en_pdf(
            tmp_path / "absent.docx", tmp_path / "sortie", executable=SOFFICE, profil=tmp_path / "profil", delai_s=120
        )


def test_conversions_simultanees_avec_un_profil_chacune(tmp_path: Path) -> None:
    # Phase 5 (T8) : plusieurs conversions en même temps, chacune avec SON profil (un profil n'en supporte qu'une).
    assert SOFFICE is not None
    from concurrent.futures import ThreadPoolExecutor

    docs = []
    for i in range(3):
        document = Document()
        document.add_paragraph(f"Document inventé numéro {i}")
        chemin = tmp_path / f"essai-{i}.docx"
        document.save(chemin)
        docs.append(chemin)
    with ThreadPoolExecutor(max_workers=3) as pool:
        pdfs = list(pool.map(lambda d: convertir_en_pdf(d, tmp_path / f"sortie-{d.stem}", executable=SOFFICE, profil=tmp_path / "profils", delai_s=120, simultanees=2), docs))
    assert [f"numéro {i}" in PdfReader(p).pages[0].extract_text() for i, p in enumerate(pdfs)] == [True] * 3
    assert sorted(p.name for p in (tmp_path / "profils").iterdir()) == ["conversion-1", "conversion-2"]


def test_pool_ne_prete_jamais_un_profil_deux_fois(tmp_path: Path) -> None:
    import threading
    import time

    en_cours: set[Path] = set()
    collisions: list[Path] = []
    verrou = threading.Lock()

    def tache() -> None:
        with _PROFILS.emprunter(tmp_path, 2) as profil:
            with verrou:
                if profil in en_cours:
                    collisions.append(profil)
                en_cours.add(profil)
                maximum.append(len(en_cours))
            time.sleep(0.02)
            with verrou:
                en_cours.discard(profil)

    maximum: list[int] = []
    fils = [threading.Thread(target=tache) for _ in range(8)]
    for f in fils:
        f.start()
    for f in fils:
        f.join()
    assert collisions == [] and max(maximum) == 2
